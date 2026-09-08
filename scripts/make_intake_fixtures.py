"""BRW-028 — deterministic synthetic intake E2E fixtures.

Builds tests/fixtures/brw024r/{sample_electrical.xlsx, sample_ultrasound.txt}:
SMALL synthetic assets matching the custom parser contracts exactly
(2 cycles, ~30 records @1s, 5 ultrasound frames × 1250 samples). These are
committed to the repo so intake E2E tests never depend on /tmp state or on
truncating the real raw assets.

Usage: .venv/bin/python scripts/make_intake_fixtures.py [--force]
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

OUT = REPO / "tests/fixtures/brw024r"

T0 = datetime(2024, 6, 1, 9, 52, 31)
RECORDS_PER_CYCLE = 15
CYCLES = 2
FRAMES = 5
SAMPLES_PER_FRAME = 1250


def make_electrical(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    unit = wb.active
    unit.title = "unit"
    unit.append(["电池号", "中文名", "化学体系", "标称容量(Ah)"])
    unit.append(["CELL_T", "测试电池", "NMC", 11.0])
    unit.append(["记录间隔(s)", 1])

    test = wb.create_sheet("test")
    test.append(["测试号", "起始时间"])
    test.append([1, T0.strftime("%Y-%m-%d %H:%M:%S")])

    cycle = wb.create_sheet("cycle")
    cycle.append(["循环号", "起始绝对时间", "结束绝对时间"])
    for c in range(1, CYCLES + 1):
        start = T0 + timedelta(seconds=(c - 1) * RECORDS_PER_CYCLE)
        end = start + timedelta(seconds=RECORDS_PER_CYCLE - 1)
        cycle.append([c, start.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S")])

    step = wb.create_sheet("step")
    step.append(["循环号", "工步号", "工步序号", "工步类型", "工步时间", "起始绝对时间", "结束绝对时间"])
    for c in range(1, CYCLES + 1):
        start = (T0 + timedelta(seconds=(c - 1) * RECORDS_PER_CYCLE)).strftime("%Y-%m-%d %H:%M:%S")
        end = (T0 + timedelta(seconds=(c - 1) * RECORDS_PER_CYCLE + RECORDS_PER_CYCLE - 1)).strftime("%Y-%m-%d %H:%M:%S")
        step.append([c, 1, 1, "恒流充电", 1, start, end])
        step.append([c, 2, 2, "恒流放电", 1, start, end])

    record = wb.create_sheet("record")
    record.append(["数据序号", "循环号", "工步号", "工步开始结束标识", "工步类型", "时间",
                   "总时间", "电流(A)", "电压(V)", "容量(Ah)", "比容量(mAh/g)",
                   "充电容量(Ah)", "充电比容量(mAh/g)", "放电容量(Ah)", "放电比容量(mAh/g)",
                   "能量(Wh)", "比能量(mWh/g)", "充电能量(Wh)", "充电比能量(mWh/g)",
                   "放电能量(Wh)", "放电比能量(mWh/g)", "绝对时间", "功率(W)",
                   "dQ/dV(mAh/V)", "dQm/dV(mAh/V.g)", "接触电阻(mΩ)", "模块启停开关",
                   "SOC/DOD(%)", "LgD"])
    idx = 0
    for c in range(1, CYCLES + 1):
        for i in range(RECORDS_PER_CYCLE):
            idx += 1
            t = T0 + timedelta(seconds=(c - 1) * RECORDS_PER_CYCLE + i)
            current = 5.0 if i < 7 else -5.0
            step_type = "恒流充电" if i < 7 else "恒流放电"
            flag = 1 if i in (0, RECORDS_PER_CYCLE - 1) else 0
            elapsed = timedelta(seconds=i + (c - 1) * RECORDS_PER_CYCLE)
            time_of_day = (T0 + timedelta(seconds=i + (c - 1) * RECORDS_PER_CYCLE)).strftime("%H:%M:%S")
            record.append([
                idx, c, 1 if i < 7 else 2, flag, step_type,
                time_of_day, str(elapsed), current,
                3.5 + 0.02 * (i % 7), 11.0 * i / RECORDS_PER_CYCLE,
                0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0,
                t.strftime("%Y-%m-%d %H:%M:%S"), 5.0 * abs(current), 0.0, 0.0, 0.0,
                1, 50.0 * (i % 7), 0,
            ])

    log = wb.create_sheet("log")
    log.append(["记录", "时间"])
    log.append(["synthetic fixture", T0.strftime("%Y-%m-%d %H:%M:%S")])
    idle = wb.create_sheet("idle")
    idle.append(["闲置"])
    idle.append([0])
    # aux sheets mirror the real contract: a unit header row then the
    # 数据序号/绝对时间/T1/… columns (validation requires those two)
    aux_vol = wb.create_sheet("auxVol")
    aux_vol.append([None, None, "单体电压(V)", None])
    aux_vol.append(["数据序号", "绝对时间", "V1", "辅助通道压差"])
    aux_temp = wb.create_sheet("auxTemp")
    aux_temp.append([None, None, "单体温度(℃)", None])
    aux_temp.append(["数据序号", "绝对时间", "T1", "辅助通道温差"])
    for i in range(1, RECORDS_PER_CYCLE * CYCLES + 1):
        t = (T0 + timedelta(seconds=i - 1)).strftime("%Y-%m-%d %H:%M:%S")
        aux_vol.append([i, t, 3.5, 0])
        aux_temp.append([i, t, 25.0, 0])
    curve = wb.create_sheet("curve")
    curve.append(["曲线"])
    curve.append([0])
    wb.save(path)


def make_ultrasound(path: Path) -> None:
    """5 frames in the EXACT custom_txt contract: 6 semicolon sections per
    line — frame_idx;unknown1;elapsed_s;meta0 meta1;waveform(1250 int32);tail(16).
    Deterministic packets: surface ~i=120, bottom ~i=850 (mirrors CELL_001)."""
    import math

    lines: list[str] = []
    for frame in range(FRAMES):
        elapsed = frame * 10.0 + 0.031217
        waveform: list[int] = []
        for i in range(SAMPLES_PER_FRAME):
            amp = 0.0
            if 100 < i < 150:
                amp = math.sin((i - 100) / 50 * math.pi * 4) * (9000 - 200 * frame)
            elif 830 < i < 900:
                amp = math.sin((i - 830) / 70 * math.pi * 4) * (5000 - 100 * frame)
            waveform.append(int(amp))
        tail = ["0"] * 16
        lines.append(
            f"{frame};0;{elapsed:.6f};27757 78;"
            + " ".join(map(str, waveform))
            + ";"
            + " ".join(tail)
        )
    path.write_text("\n".join(lines) + "\n", encoding="ascii")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx = OUT / "sample_electrical.xlsx"
    txt = OUT / "sample_ultrasound.txt"
    if xlsx.is_file() and txt.is_file() and not args.force:
        print("fixtures already present:", OUT)
        return 0
    make_electrical(xlsx)
    make_ultrasound(txt)
    print("fixtures written →", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
